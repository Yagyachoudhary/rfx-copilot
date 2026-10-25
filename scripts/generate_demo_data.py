"""
Generates the fictional demo dataset for RFx Copilot.

Everything here is FICTIONAL: buyer, suppliers, people, prices, certificates.
The app never reads the price tables below. It only sees the rendered vendor
documents (xlsx / pdf / docx / jpg / email) and must extract them at runtime.

Run:  python3 scripts/generate_demo_data.py
"""
import json, math, os, random
from datetime import date

random.seed(142)
ROOT = os.path.join(os.path.dirname(__file__), "..", "sample-data")
RESP = os.path.join(ROOT, "responses")
os.makedirs(RESP, exist_ok=True)

USD_INR = 84.20

# ----------------------------------------------------------------------------
# Buyer RFQ: 30 lines of corrugated packaging
# (line_no, description, specification, qty, uom, category, fair_price_inr)
# ----------------------------------------------------------------------------
LINES = [
    (1, "3-ply RSC carton 250x200x150 mm", "3-ply, B-flute, 150 GSM kraft liner, plain, ID dimensions", 40000, "pcs", "3-ply cartons", 9.20),
    (2, "3-ply RSC carton 300x200x150 mm, 1-colour print", "3-ply, B-flute, 1-colour flexo print (logo), ID dimensions", 60000, "pcs", "3-ply cartons", 11.40),
    (3, "3-ply RSC carton 350x250x200 mm", "3-ply, B-flute, 150 GSM kraft, plain", 35000, "pcs", "3-ply cartons", 15.80),
    (4, "3-ply RSC carton 400x300x200 mm", "3-ply, B-flute, 150 GSM kraft, plain", 25000, "pcs", "3-ply cartons", 19.60),
    (5, "3-ply RSC carton 200x150x100 mm (e-com small)", "3-ply, B-flute, 120 GSM kraft, plain", 50000, "pcs", "3-ply cartons", 5.40),
    (6, "3-ply RSC carton 450x350x250 mm", "3-ply, C-flute, 180 GSM kraft, plain", 15000, "pcs", "3-ply cartons", 26.50),
    (7, "5-ply RSC carton 400x300x300 mm", "5-ply, BC-flute, BCT >= 350 kgf, plain", 30000, "pcs", "5-ply cartons", 31.00),
    (8, "5-ply RSC master carton 450x300x300 mm", "5-ply, BC-flute, BCT >= 400 kgf, 1-colour print", 25000, "pcs", "5-ply cartons", 36.50),
    (9, "5-ply RSC carton 500x400x300 mm", "5-ply, BC-flute, BCT >= 420 kgf, plain", 18000, "pcs", "5-ply cartons", 44.00),
    (10, "5-ply RSC carton 600x400x400 mm", "5-ply, BC-flute, BCT >= 500 kgf, plain", 12000, "pcs", "5-ply cartons", 62.00),
    (11, "5-ply RSC carton 350x300x250 mm, 2-colour print", "5-ply, BC-flute, 2-colour flexo print", 20000, "pcs", "5-ply cartons", 30.50),
    (12, "5-ply RSC carton 550x450x350 mm", "5-ply, BC-flute, BCT >= 520 kgf, plain", 8000, "pcs", "5-ply cartons", 58.00),
    (13, "7-ply heavy-duty carton 600x500x500 mm", "7-ply, BCB-flute, BCT >= 900 kgf", 4000, "pcs", "7-ply cartons", 118.00),
    (14, "7-ply heavy-duty export carton 450x300x300 mm", "7-ply, BCB-flute, export grade, BCT >= 750 kgf", 5000, "pcs", "7-ply cartons", 52.00),
    (15, "Die-cut mailer box 300x220x80 mm, 2-colour print", "E-flute, die-cut, tuck-front, 2-colour print", 30000, "pcs", "Mailers & trays", 14.20),
    (16, "Die-cut mailer box 220x160x60 mm", "E-flute, die-cut, tuck-front, plain", 40000, "pcs", "Mailers & trays", 8.10),
    (17, "Die-cut tray with lid 400x300x100 mm", "3-ply, B-flute, two-piece tray + lid", 10000, "pcs", "Mailers & trays", 21.50),
    (18, "Corrugated partition set, 12-cell (1 L bottles)", "3-ply, B-flute, interlocking, fits line 7 carton", 20000, "sets", "Partitions & pads", 7.80),
    (19, "Corrugated partition set, 24-cell", "3-ply, B-flute, interlocking, fits line 9 carton", 10000, "sets", "Partitions & pads", 12.60),
    (20, "Corrugated partition set, 6-cell", "3-ply, B-flute, interlocking, fits line 3 carton", 15000, "sets", "Partitions & pads", 4.90),
    (21, "Corrugated layer pad 1200x1000 mm", "3-ply, B-flute, pallet layer pad", 12000, "pcs", "Partitions & pads", 28.00),
    (22, "Corrugated sheet roll 1 m x 50 m", "2-ply single-face, E-flute", 400, "rolls", "Partitions & pads", 640.00),
    (23, "Paper edge protector 50x50x5 mm x 1000 mm", "Laminated kraft, L-profile", 25000, "pcs", "Protective", 9.50),
    (24, "Paper edge protector 35x35x3 mm x 800 mm", "Laminated kraft, L-profile", 20000, "pcs", "Protective", 5.20),
    (25, "BOPP tape 48 mm x 65 m, clear", "40 micron, acrylic adhesive", 18000, "rolls", "Tapes & films", 34.00),
    (26, "BOPP tape 72 mm x 65 m, printed logo", "45 micron, 1-colour logo print", 6000, "rolls", "Tapes & films", 58.00),
    (27, "Stretch film 500 mm x 23 micron, hand roll", "LLDPE cast, hand grade", 6000, "kg", "Tapes & films", 168.00),
    (28, "Stretch film 500 mm x 17 micron, machine grade", "LLDPE cast, 250% pre-stretch", 4000, "kg", "Tapes & films", 176.00),
    (29, "Corrugated pallet sleeve 1200x1000x1000 mm", "5-ply, BC-flute, with top cap", 1500, "pcs", "Protective", 410.00),
    (30, "Kraft paper void-fill roll 450 mm x 300 m", "70 GSM recycled kraft", 800, "rolls", "Protective", 1150.00),
]
FAIR = {l[0]: l[6] for l in LINES}

QUESTIONS = [
    (1, "Do you hold a valid ISO 9001:2015 certificate? Please attach it.", "yes_no", True, {"op": "eq", "value": True}, True),
    (2, "Can you provide BCT/ECT test reports for every batch (NABL lab or calibrated in-house lab)?", "yes_no", True, {"op": "eq", "value": True}, False),
    (3, "Monthly corrugation capacity you can commit to this contract (tonnes/month)?", "number", True, {"op": "gte", "value": 150}, False),
    (4, "Standard lead time from PO to delivery at Hosur DC for regular cartons (days)?", "number", True, {"op": "lte", "value": 14}, False),
    (5, "Do you accept 60-day payment terms from invoice date?", "yes_no", False, None, False),
    (6, "Can you supply FSC-certified or 100% recycled-fibre board on request?", "yes_no", False, None, False),
    (7, "Distance of your manufacturing plant from Hosur DC (km)?", "number", False, None, False),
    (8, "Name two current FMCG / food customers we may contact as references.", "text", False, None, False),
]

SUPPLIERS = [
    {"key": "shree-balaji", "name": "Shree Balaji Corrugators", "short_code": "SBC", "contact_name": "Venkatesh R.", "email": "sales@shreebalaji-demo.example", "phone": "+91 80 0000 1101", "location": "Peenya Industrial Area, Bengaluru"},
    {"key": "deccan", "name": "Deccan Packaging Industries Pvt Ltd", "short_code": "DPI", "contact_name": "Anita Rao", "email": "quotes@deccanpack-demo.example", "phone": "+91 4344 000 220", "location": "SIPCOT Phase II, Hosur"},
    {"key": "pacific", "name": "Pacific Pack Solutions Pte Ltd", "short_code": "PPS", "contact_name": "Marcus Tan", "email": "marcus.tan@pacificpack-demo.example", "phone": "+65 6000 0330", "location": "Singapore (warehouse: Sriperumbudur, Chennai)"},
    {"key": "hosur-box", "name": "Hosur Box Works", "short_code": "HBW", "contact_name": "Murugan S.", "email": "hosurboxworks-demo@example.com", "phone": "+91 98 0000 0440", "location": "Bagalur Road, Hosur"},
    {"key": "sri-lakshmi", "name": "Sri Lakshmi Paper Products", "short_code": "SLP", "contact_name": "Ramesh K.", "email": "ramesh.slpp-demo@example.com", "phone": "+91 99 0000 0550", "location": "Jigani, Bengaluru"},
]


def r2(x):
    return round(x + 1e-9, 2)


def vendor_price(line_no, factor, jitter=0.04):
    return FAIR[line_no] * factor * (1 + random.uniform(-jitter, jitter))


# ============================================================================
# 1) Shree Balaji — Excel that ignores the template (v1 has a typo + duplicate)
# ============================================================================
def make_shree_balaji():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    prices = {n: vendor_price(n, 1.03) for n in FAIR}
    # quoted per 100 Nos for small items
    per100 = {5, 16, 23, 24}
    order = [1, 2, 5, 3, 4, 6, 7, 8, 11, 9, 10, 12, 13, 14, 15, 16, 17, 20, 18, 19, 21, 22, 23, 24, 25, 26, 27, 28, 29, 30]
    descr = {
        1: ("3 Ply Box", "250 x 200 x 150", "3", "Plain"),
        2: ("3 Ply Box Printed 1 col", "300 x 200 x 150", "3", "Logo print"),
        3: ("3 Ply Box", "350 x 250 x 200", "3", "Plain"),
        4: ("3 Ply Box", "400 x 300 x 200", "3", "Plain"),
        5: ("3 Ply Box (small)", "200 x 150 x 100", "3", "Rate per 100 Nos"),
        6: ("3 Ply Box C-flute", "450 x 350 x 250", "3", "180 GSM"),
        7: ("5 Ply Box", "400 x 300 x 300", "5", "BCT 350"),
        8: ("5 Ply Master Carton Printed", "450 x 300 x 300", "5", "BCT 400"),
        9: ("5 Ply Box", "500 x 400 x 300", "5", ""),
        10: ("5 Ply Box", "600 x 400 x 400", "5", ""),
        11: ("5 Ply Box 2 col print", "350 x 300 x 250", "5", ""),
        12: ("5 Ply Box", "550 x 450 x 350", "5", ""),
        13: ("7 Ply HD Box", "600 x 500 x 500", "7", "BCT 900"),
        14: ("7 Ply Export Box", "450 x 300 x 300", "7", "Export grade"),
        15: ("Mailer Box E-flute 2 col", "300 x 220 x 80", "E", "Die cut"),
        16: ("Mailer Box E-flute", "220 x 160 x 60", "E", "Rate per 100 Nos"),
        17: ("Tray + Lid", "400 x 300 x 100", "3", "2 pc"),
        18: ("Partition 12 cell", "for 400x300x300", "3", "per set"),
        19: ("Partition 24 cell", "for 500x400x300", "3", "per set"),
        20: ("Partition 6 cell", "for 350x250x200", "3", "per set"),
        21: ("Layer Pad", "1200 x 1000", "3", ""),
        22: ("Corrugated Roll 2 ply", "1 mtr x 50 mtr", "2", "per roll"),
        23: ("Edge Board 50x50x5", "1000 mm", "-", "Rate per 100 Nos"),
        24: ("Edge Board 35x35x3", "800 mm", "-", "Rate per 100 Nos"),
        25: ("BOPP Tape Clear 48mm", "65 mtr", "-", "40 mic, per roll"),
        26: ("BOPP Tape Printed 72mm", "65 mtr", "-", "45 mic, per roll"),
        27: ("Stretch Film Hand 23 mic", "500 mm", "-", "per Kg"),
        28: ("Stretch Film Machine 17 mic", "500 mm", "-", "per Kg"),
        29: ("Pallet Sleeve with cap", "1200 x 1000 x 1000", "5", ""),
        30: ("Kraft Void Fill Roll 70 GSM", "450 mm x 300 m", "-", "per roll"),
    }

    def build(version):
        wb = Workbook()
        ws = wb.active
        ws.title = "Quotation"
        bold = Font(bold=True)
        ws["A1"] = "SHREE BALAJI CORRUGATORS"
        ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
        ws["A2"] = "Plot 42, Peenya Industrial Area, Bengaluru 560058  |  GSTIN 29XXXXX0000X1Z0 (demo)"
        ws["A3"] = f"Quotation No: SBC/Q/26-27/{'0418' if version == 1 else '0418-R1'}    Date: {'29-10-2026' if version == 1 else '02-11-2026'}"
        ws["A4"] = "Kind Attn: Procurement Team, Nilgiri Consumer Foods (DEMO)   Ref: Your enquiry RFQ-2026-0142"
        if version == 2:
            ws["A5"] = "REVISED QUOTATION - supersedes SBC/Q/26-27/0418. Item 3 rate corrected, duplicate line removed."
            ws["A5"].font = Font(bold=True, color="C00000")
        hdr = ["Sl", "Our Code", "Item", "Size (mm)", "Ply", "Rate (Rs)", "Per", "Remarks"]
        r0 = 7
        fill = PatternFill("solid", fgColor="DDEBF7")
        thin = Side(style="thin", color="999999")
        for i, h in enumerate(hdr):
            c = ws.cell(row=r0, column=i + 1, value=h)
            c.font = bold
            c.fill = fill
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        row = r0 + 1
        sl = 1
        for n in order:
            item, size, ply, rem = descr[n]
            p = prices[n]
            per = "Nos"
            if n in per100:
                p = p * 100
                per = "100 Nos"
            if n in (22, 25, 26, 30):
                per = "Roll"
            if n in (27, 28):
                per = "Kg"
            if n in (18, 19, 20):
                per = "Set"
            rate = r2(p)
            if version == 1 and n == 3:
                rate = r2(prices[3] / 10)  # decimal typo: 1.6x instead of 16.x
            if n == 8:
                rate = 41.50 if version == 1 else 39.80
            code = f"SBC-{ply if ply not in '-E' else 'X'}P-{n:03d}" if ply != "-" else f"SBC-ACC-{n:03d}"
            vals = [sl, code, item, size, ply, rate, per, rem]
            for i, v in enumerate(vals):
                ws.cell(row=row, column=i + 1, value=v).border = Border(top=thin, bottom=thin, left=thin, right=thin)
            row += 1
            sl += 1
        if version == 1:
            # duplicate entry for line 8 with a different rate -> conflict
            vals = [sl, "SBC-5P-008A", "5 Ply Master Carton Printed (revised)", "450 x 300 x 300", "5", 39.80, "Nos", "revised rate"]
            for i, v in enumerate(vals):
                ws.cell(row=row, column=i + 1, value=v)
            row += 1
        row += 1
        ws.cell(row=row, column=1, value="Note: Rates exclusive of GST @ 18%. Freight included up to Hosur DC. See 'Terms' sheet.")
        for col, w in zip("ABCDEFGH", [5, 14, 34, 20, 6, 11, 9, 22]):
            ws.column_dimensions[col].width = w

        t = wb.create_sheet("Terms")
        terms = [
            ("Taxes", "GST @ 18% extra as applicable"),
            ("Freight", "Included - door delivery Hosur DC"),
            ("Payment", "45 days from invoice"),
            ("Validity", "30 days from quotation date"),
            ("Delivery", "10 working days from PO for regular items; 15 days for printed items"),
            ("MOQ", "1,000 Nos per size for printed items"),
        ]
        t["A1"] = "TERMS & CONDITIONS"
        t["A1"].font = bold
        for i, (k, v) in enumerate(terms):
            t.cell(row=i + 3, column=1, value=k).font = bold
            t.cell(row=i + 3, column=2, value=v)
        t.column_dimensions["A"].width = 14
        t.column_dimensions["B"].width = 70

        q = wb.create_sheet("Vendor Questionnaire")
        q["A1"] = "SUPPLIER QUESTIONNAIRE - RESPONSES"
        q["A1"].font = bold
        answers = [
            ("ISO 9001:2015 certified?", "Yes - certificate available on request"),
            ("BCT / ECT test reports per batch?", "Yes, in-house calibrated BCT tester"),
            ("Monthly capacity for this contract (MT)", "220 MT / month"),
            ("Lead time (days)", "10 days regular, 15 days printed"),
            ("60 days payment acceptable?", "No - our standard is 45 days"),
            ("FSC / recycled board?", "Yes, FSC Mix available"),
            ("Plant distance from Hosur DC", "Approx. 38 km"),
            ("Two FMCG references", "Kaveri Snacks (demo), Malnad Dairy (demo)"),
        ]
        for i, (k, v) in enumerate(answers):
            q.cell(row=i + 3, column=1, value=f"Q{i+1}. {k}")
            q.cell(row=i + 3, column=2, value=v)
        q.column_dimensions["A"].width = 44
        q.column_dimensions["B"].width = 50
        d = os.path.join(RESP, "shree-balaji")
        os.makedirs(d, exist_ok=True)
        name = "SBC_Quotation_RFQ-0142.xlsx" if version == 1 else "SBC_Quotation_RFQ-0142_REVISED.xlsx"
        wb.save(os.path.join(d, name))
        return name

    return build(1), build(2)


# ============================================================================
# 2) Deccan — PDF on letterhead, 27 of 30 lines, discount + freight in footnotes
# ============================================================================
def make_deccan():
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    d = os.path.join(RESP, "deccan")
    os.makedirs(d, exist_ok=True)
    path = os.path.join(d, "DPI_Quotation_RFQ-2026-0142.pdf")
    styles = getSampleStyleSheet()
    small = ParagraphStyle("small", parent=styles["Normal"], fontSize=7, leading=9, textColor=colors.HexColor("#444444"))
    norm = ParagraphStyle("n", parent=styles["Normal"], fontSize=9, leading=12)

    def letterhead(c, doc):
        c.saveState()
        c.setFillColor(colors.HexColor("#0B5D3B"))
        c.rect(0, A4[1] - 28 * mm, A4[0], 28 * mm, fill=1, stroke=0)
        c.setFillColor(colors.white)
        c.setFont("Helvetica-Bold", 17)
        c.drawString(18 * mm, A4[1] - 14 * mm, "DECCAN PACKAGING INDUSTRIES PVT LTD")
        c.setFont("Helvetica", 8.5)
        c.drawString(18 * mm, A4[1] - 20 * mm, "Corrugated boxes | Partitions | Protective packaging   -   SIPCOT Phase II, Hosur 635109 (DEMO)")
        c.drawString(18 * mm, A4[1] - 24.5 * mm, "ISO 9001:2015  |  quotes@deccanpack-demo.example  |  +91 4344 000 220")
        c.setFillColor(colors.HexColor("#666666"))
        c.setFont("Helvetica", 7)
        c.drawString(18 * mm, 10 * mm, f"Fictional document generated for a product prototype demo.   Page {doc.page}")
        c.restoreState()

    doc = SimpleDocTemplate(path, pagesize=A4, topMargin=34 * mm, leftMargin=16 * mm, rightMargin=16 * mm, bottomMargin=18 * mm)
    story = []
    story.append(Paragraph("<b>QUOTATION</b> &nbsp;&nbsp; No. DPI/HSR/2026/1187 &nbsp;&nbsp; Date: 30 Oct 2026", norm))
    story.append(Paragraph("To: Procurement - Packaging, Nilgiri Consumer Foods (DEMO), Hosur DC. Ref: RFQ-2026-0142.", norm))
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph("We thank you for your enquiry and are pleased to quote our best rates as under:", norm))
    story.append(Spacer(1, 3 * mm))

    missing = {22, 27, 28}
    data = [["S.No", "Description", "Size / Spec", "Qty", "Unit", "Rate*\n(Rs/unit)", "Amount (Rs)"]]
    sno = 1
    total = 0.0
    for (n, desc, spec, qty, uom, cat, fair) in LINES:
        if n in missing:
            continue
        rate = r2(vendor_price(n, 0.98))
        amt = rate * qty
        total += amt
        short = desc.replace(" mm", "").replace("RSC ", "")
        data.append([str(sno), short, spec[:44], f"{qty:,}", uom, f"{rate:,.2f}", f"{amt:,.2f}"])
        sno += 1
    data.append(["", "", "", "", "", "Basic total", f"{total:,.2f}"])
    t = Table(data, colWidths=[10 * mm, 62 * mm, 50 * mm, 16 * mm, 11 * mm, 17 * mm, 22 * mm], repeatRows=1)
    t.setStyle(TableStyle([
        ("FONT", (0, 0), (-1, -1), "Helvetica", 7.2),
        ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 7.2),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E2EFE8")),
        ("GRID", (0, 0), (-1, -2), 0.3, colors.HexColor("#999999")),
        ("ALIGN", (3, 1), (-1, -1), "RIGHT"),
        ("FONT", (5, -1), (-1, -1), "Helvetica-Bold", 7.5),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story.append(t)
    story.append(Spacer(1, 3 * mm))
    story.append(Paragraph("Note: We regret we are unable to quote for the corrugated sheet roll and stretch film items at this time.", norm))
    story.append(Spacer(1, 5 * mm))
    story.append(Paragraph("<b>Terms</b>: GST @ 18% extra. Payment: 60 days from invoice acceptable. Validity: 60 days. Delivery: 7 working days from PO.", norm))
    story.append(Spacer(1, 30 * mm))
    story.append(Paragraph("* Rates are ex-works Hosur. Freight to customer DC will be charged extra @ 2.5% of basic value.", small))
    story.append(Paragraph("** A special discount of 3% on basic rates is applicable for purchase orders released on or before 30-Nov-2026.", small))
    story.append(Paragraph("For Deccan Packaging Industries Pvt Ltd - Anita Rao, Sales Head (signature on file)", small))

    from reportlab.platypus import PageBreak
    story.append(PageBreak())
    story.append(Paragraph("<b>ANNEXURE A - Supplier questionnaire</b>", norm))
    story.append(Spacer(1, 3 * mm))
    qa = [
        ["Q1", "ISO 9001:2015 certificate", "Yes. Certificate no. QMS-DEMO-44817 attached (valid till 14-Mar-2027)."],
        ["Q2", "BCT/ECT reports per batch", "Yes - in-house lab, equipment calibrated annually (NABL-traceable)."],
        ["Q3", "Monthly capacity committed", "300 tonnes per month"],
        ["Q4", "Lead time to Hosur DC", "7 working days"],
        ["Q5", "60-day payment terms", "Accepted"],
        ["Q6", "FSC / recycled board", "Recycled board standard; FSC on request at +4%"],
        ["Q7", "Plant distance from Hosur DC", "6 km"],
        ["Q8", "References", "Cauvery Beverages (demo); Western Ghats Foods (demo)"],
    ]
    t2 = Table(qa, colWidths=[10 * mm, 55 * mm, 115 * mm])
    t2.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", 8), ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#999999")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story.append(t2)
    doc.build(story, onFirstPage=letterhead, onLaterPages=letterhead)

    # Attachment: ISO certificate (fictional)
    from reportlab.pdfgen import canvas
    cpath = os.path.join(d, "DPI_ISO9001_Certificate.pdf")
    c = canvas.Canvas(cpath, pagesize=A4)
    W, H = A4
    c.setStrokeColor(colors.HexColor("#8A6D3B"))
    c.setLineWidth(4)
    c.rect(15 * mm, 15 * mm, W - 30 * mm, H - 30 * mm)
    c.setFont("Helvetica-Bold", 22)
    c.drawCentredString(W / 2, H - 60 * mm, "CERTIFICATE OF REGISTRATION")
    c.setFont("Helvetica", 12)
    c.drawCentredString(W / 2, H - 75 * mm, "This is to certify that the Quality Management System of")
    c.setFont("Helvetica-Bold", 16)
    c.drawCentredString(W / 2, H - 90 * mm, "Deccan Packaging Industries Pvt Ltd")
    c.setFont("Helvetica", 11)
    c.drawCentredString(W / 2, H - 100 * mm, "SIPCOT Phase II, Hosur 635109, Tamil Nadu, India")
    c.drawCentredString(W / 2, H - 115 * mm, "has been assessed and found to conform to the requirements of")
    c.setFont("Helvetica-Bold", 14)
    c.drawCentredString(W / 2, H - 128 * mm, "ISO 9001:2015")
    c.setFont("Helvetica", 11)
    c.drawCentredString(W / 2, H - 143 * mm, "Scope: Manufacture and supply of corrugated boxes, partitions and protective packaging")
    c.drawCentredString(W / 2, H - 165 * mm, "Certificate No: QMS-DEMO-44817     Issued: 15-Mar-2024     Valid until: 14-Mar-2027")
    c.setFont("Helvetica-Oblique", 9)
    c.setFillColor(colors.HexColor("#B00020"))
    c.drawCentredString(W / 2, 30 * mm, "SPECIMEN - FICTIONAL CERTIFICATE FOR PROTOTYPE DEMO - NOT ISSUED BY ANY CERTIFICATION BODY")
    c.save()
    return os.path.basename(path), os.path.basename(cpath)


# ============================================================================
# 3) Pacific — Word letter, commercials in paragraphs, USD, per-box tapes
# ============================================================================
def make_pacific():
    from docx import Document
    from docx.shared import Pt

    p = {n: vendor_price(n, 0.93) / USD_INR for n in FAIR}
    d = os.path.join(RESP, "pacific")
    os.makedirs(d, exist_ok=True)
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    doc.add_heading("Pacific Pack Solutions Pte Ltd", level=1)
    doc.add_paragraph("80 Robinson Road, #10-01, Singapore 068898  |  India warehouse: Sriperumbudur, Chennai  (DEMO COMPANY)")
    doc.add_paragraph("31 October 2026")
    doc.add_paragraph("To: Nilgiri Consumer Foods (DEMO) - Procurement\nSubject: Our offer against your RFQ-2026-0142 (corrugated packaging, FY27 H1)")
    doc.add_paragraph("Dear Procurement Team,")
    doc.add_paragraph(
        "Thank you for inviting Pacific Pack Solutions to participate. As a regional converter with a bonded warehouse in "
        "Chennai, we are pleased to offer the following. All prices below are in US Dollars (USD), delivered to your Hosur DC, "
        "inclusive of import duties but exclusive of Indian GST.")
    doc.add_heading("Corrugated cartons", level=2)
    doc.add_paragraph(
        f"For the 3-ply range our prices are USD {p[1]:.3f} per piece for 250x200x150 mm, USD {p[2]:.3f} for the 300x200x150 mm "
        f"printed carton, USD {p[3]:.3f} for 350x250x200 mm and USD {p[4]:.3f} for 400x300x200 mm. The small e-commerce carton "
        f"(200x150x100 mm) is USD {p[5]*1000:.2f} per thousand pieces, and the 450x350x250 mm C-flute carton is USD {p[6]:.3f} per piece.")
    doc.add_paragraph(
        f"In 5-ply we offer 400x300x300 mm at USD {p[7]:.3f}, 500x400x300 mm at USD {p[9]:.3f}, 600x400x400 mm at USD {p[10]:.3f}, "
        f"the 2-colour 350x300x250 mm at USD {p[11]:.3f} and 550x450x350 mm at USD {p[12]:.3f}, all per piece.")
    doc.add_paragraph(
        f"For the heavy-duty 450x300x300 mm carton we can offer USD {p[8]*1.06:.3f} per piece. Our 7-ply 600x500x500 mm "
        f"heavy-duty carton is USD {p[13]:.3f} per piece.")
    doc.add_heading("Mailers, partitions and pads", level=2)
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "Item", "Price (USD)", "Basis"
    for n, name, basis in [
        (15, "Mailer 300x220x80 E-flute, 2-col", "per pc"),
        (16, "Mailer 220x160x60 E-flute", "per pc"),
        (17, "Tray + lid 400x300x100", "per pc"),
        (18, "Partition 12-cell", "per set"),
        (19, "Partition 24-cell", "per set"),
        (20, "Partition 6-cell", "per set"),
        (21, "Layer pad 1200x1000, 3-ply", "per pc"),
        (22, "Single-face roll 1 m x 50 m", "per roll"),
    ]:
        row = t.add_row().cells
        row[0].text, row[1].text, row[2].text = name, f"{p[n]:.3f}" if p[n] < 5 else f"{p[n]:.2f}", basis
    doc.add_heading("Tapes and films", level=2)
    doc.add_paragraph(
        f"Clear BOPP tape 48 mm x 65 m (40 mic) is USD {p[25]*36:.2f} per box; each carton box contains 36 rolls. "
        f"Printed 72 mm x 65 m logo tape is USD {p[26]*24:.2f} per box. "
        f"Hand stretch film 23 mic is USD {p[27]:.2f}/kg and machine film 17 mic is USD {p[28]:.2f}/kg.")
    doc.add_paragraph(
        "We are not quoting edge protectors, pallet sleeves or void-fill rolls in this round as these are sourced locally "
        "by our partners; we would be happy to include them in a second round.")
    doc.add_heading("Commercial terms", level=2)
    doc.add_paragraph(
        "Payment: 30 days from bill of lading against irrevocable LC. Validity: 45 days. Lead time: 21 days from PO "
        "(ex-Singapore converting, cleared and delivered to Hosur). Exchange rate risk is with the buyer; invoices will be in USD.")
    doc.add_heading("Supplier questionnaire", level=2)
    for qline in [
        "1. ISO 9001:2015 - Yes, our Singapore plant is ISO 9001:2015 certified (certificate to follow).",
        "2. BCT/ECT reports - Yes, third-party SGS reports per shipment.",
        "3. Capacity - 500 tonnes per month across our regional network.",
        "4. Lead time - 21 days to Hosur DC.",
        "5. 60-day payment - Not possible; LC 30 days.",
        "6. FSC - Yes, FSC-certified board available as standard.",
        "7. Distance - Chennai warehouse is approx. 280 km from Hosur.",
        "8. References - Straits Confectionery (demo); Lion City Foods (demo).",
    ]:
        doc.add_paragraph(qline)
    doc.add_paragraph("Warm regards,\nMarcus Tan, Regional Sales Director\nThis is a fictional demo document.")
    name = "Pacific_Pack_Offer_RFQ-2026-0142.docx"
    doc.save(os.path.join(d, name))
    return name


# ============================================================================
# 4) Hosur Box Works — printed rate card photographed at an angle
# ============================================================================
def make_hosur_photo():
    import numpy as np
    import cv2
    from PIL import Image, ImageDraw, ImageFont, ImageFilter

    d = os.path.join(RESP, "hosur-box")
    os.makedirs(d, exist_ok=True)
    W, H = 1500, 2100
    img = Image.new("RGB", (W, H), (246, 243, 234))
    dr = ImageDraw.Draw(img)

    def font(sz, bold=False):
        for f in (["/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"] if bold else ["/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"]):
            if os.path.exists(f):
                return ImageFont.truetype(f, sz)
        return ImageFont.load_default()

    dr.text((80, 70), "HOSUR BOX WORKS", fill=(20, 20, 90), font=font(62, True))
    dr.text((80, 150), "Bagalur Road, Hosur  -  Ph 98 0000 0440  (DEMO)", fill=(40, 40, 40), font=font(30))
    dr.text((80, 215), "RATE CARD - effective 01 Oct 2026", fill=(150, 20, 20), font=font(40, True))
    dr.text((80, 270), "All rates in Rupees, INCLUSIVE of GST 18%, delivered Hosur", fill=(40, 40, 40), font=font(30))
    rows = [
        ("3 Ply Box 250x200x150", 1, "pc"), ("3 Ply Box 300x200x150 (1 col)", 2, "pc"), ("3 Ply Box 350x250x200", 3, "pc"),
        ("3 Ply Box 400x300x200", 4, "pc"), ("3 Ply Box 200x150x100", 5, "pc"), ("5 Ply Box 400x300x300", 7, "pc"),
        ("5 Ply Box 450x300x300", 8, "pc"), ("5 Ply Box 500x400x300", 9, "pc"), ("5 Ply Box 600x400x400", 10, "pc"),
        ("5 Ply Box 350x300x250 (2 col)", 11, "pc"), ("7 Ply Box 600x500x500", 13, "pc"), ("Mailer E-flute 220x160x60", 16, "pc"),
        ("Partition 12 cell", 18, "bundle of 50 sets"), ("Partition 6 cell", 20, "bundle of 50 sets"), ("Layer Pad 1200x1000", 21, "pc"),
        ("Edge Protector 50x50x5 1m", 23, "100 pcs"), ("Edge Protector 35x35x3 0.8m", 24, "100 pcs"), ("Pallet Sleeve 1200x1000x1000", 29, "pc"),
    ]
    y = 360
    dr.line((70, y - 10, W - 70, y - 10), fill=(30, 30, 30), width=3)
    dr.text((80, y), "ITEM", fill=(0, 0, 0), font=font(32, True))
    dr.text((900, y), "RATE", fill=(0, 0, 0), font=font(32, True))
    dr.text((1120, y), "PER", fill=(0, 0, 0), font=font(32, True))
    y += 60
    dr.line((70, y - 10, W - 70, y - 10), fill=(30, 30, 30), width=2)
    for label, n, per in rows:
        base = vendor_price(n, 0.97) * 1.18
        mult = 50 if per.startswith("bundle") else 100 if per == "100 pcs" else 1
        val = round(base * mult, 0 if mult > 1 or base > 100 else 1)
        dr.text((80, y), label, fill=(25, 25, 25), font=font(34))
        dr.text((900, y), f"{val:,.1f}" if val < 100 else f"{val:,.0f}", fill=(25, 25, 25), font=font(34, True))
        dr.text((1120, y), per, fill=(25, 25, 25), font=font(30))
        y += 82
        dr.line((70, y - 14, W - 70, y - 14), fill=(170, 170, 170), width=1)
    dr.text((80, y + 30), "Printing extra @ Rs 0.80 per colour per box. Rates subject to paper price change.", fill=(60, 60, 60), font=font(28))
    dr.text((80, y + 80), "Fictional rate card generated for a prototype demo.", fill=(120, 120, 120), font=font(24))

    card = np.array(img)
    # perspective warp onto a desk background
    BW, BH = 1700, 2200
    bg = np.zeros((BH, BW, 3), np.uint8)
    yy, xx = np.mgrid[0:BH, 0:BW]
    bg[..., 0] = (95 + 25 * np.sin(xx / 37.0) + 10 * np.sin(yy / 11.0)).astype(np.uint8)
    bg[..., 1] = (70 + 18 * np.sin(xx / 37.0)).astype(np.uint8)
    bg[..., 2] = (52 + 12 * np.sin(xx / 37.0)).astype(np.uint8)
    src = np.float32([[0, 0], [W, 0], [W, H], [0, H]])
    dst = np.float32([[290, 190], [1430, 110], [1600, 2110], [70, 2080]])
    M = cv2.getPerspectiveTransform(src, dst)
    warped = cv2.warpPerspective(cv2.cvtColor(card, cv2.COLOR_RGB2BGR), M, (BW, BH))
    mask = cv2.warpPerspective(np.ones((H, W), np.uint8) * 255, M, (BW, BH))
    bgb = cv2.cvtColor(bg, cv2.COLOR_RGB2BGR)
    out = np.where(mask[..., None] > 0, warped, bgb).astype(np.float32)
    # diagonal shadow (phone/hand) across lower-right
    shade = np.ones((BH, BW), np.float32)
    shade -= 0.55 * np.clip((xx + yy - 2300) / 900.0, 0, 1)
    out *= shade[..., None]
    # global slight blur + strong blur on the lower block (unreadable prices)
    out = cv2.GaussianBlur(out, (3, 3), 0)
    lower = out[1500:2050, 700:1500]
    out[1500:2050, 700:1500] = cv2.GaussianBlur(lower, (0, 0), 7.5)
    noise = np.random.normal(0, 6, out.shape).astype(np.float32)
    out = np.clip(out + noise, 0, 255).astype(np.uint8)
    # slight rotation
    R = cv2.getRotationMatrix2D((BW / 2, BH / 2), -1.0, 1.0)
    out = cv2.warpAffine(out, R, (BW, BH), borderMode=cv2.BORDER_REFLECT)
    name = "IMG_20261029_174512_ratecard.jpg"
    cv2.imwrite(os.path.join(d, name), out, [cv2.IMWRITE_JPEG_QUALITY, 72])
    return name


# ============================================================================
# 5) Sri Lakshmi — email body with partial pricing
# ============================================================================
def make_sri_lakshmi():
    d = os.path.join(RESP, "sri-lakshmi")
    os.makedirs(d, exist_ok=True)
    body = """From: Ramesh K. <ramesh.slpp-demo@example.com>
To: packaging.buying@nilgiri-demo.example
Date: Mon, 2 Nov 2026 09:14:00 +0530
Subject: RE: RFQ-2026-0142 Corrugated packaging - quote

Dear Sir,

Thanks for the enquiry. Rs 42/kg for the 5-ply, 38 for the 3-ply, rest same as last year, freight extra.
48mm tape we can do at Rs 31 per roll, 72mm printed tape Rs 55.

Material will be as per your spec. Delivery 2 weeks from PO. Payment 60 days ok for us.

Regards,
Ramesh
Sri Lakshmi Paper Products, Jigani
(Fictional email for a prototype demo)
"""
    name = "RE_RFQ-2026-0142_quote.eml"
    with open(os.path.join(d, name), "w") as f:
        f.write(body)
    return name


def main():
    sb1, sb2 = make_shree_balaji()
    dpi, dpi_cert = make_deccan()
    pps = make_pacific()
    hbw = make_hosur_photo()
    slp = make_sri_lakshmi()

    manifest = {
        "_note": "Fictional demo data. Buyer, suppliers, people, certificates and prices are invented.",
        "buyer_org_name": "Nilgiri Consumer Foods (Demo)",
        "rfq": {
            "ref_code": "RFQ-2026-0142",
            "title": "Corrugated packaging - FY27 H1 rate contract (Hosur DC)",
            "description": "Six-month rate contract for corrugated cartons, mailers, partitions, protective packaging, tapes and stretch film for the Hosur distribution centre. Prices to be delivered Hosur DC, GST shown separately.",
            "category": "Corrugated packaging",
            "currency": "INR",
            "response_deadline": "2026-11-06",
            "required_by": "2026-12-01",
            "delivery_location": "Hosur DC, Tamil Nadu (demo)",
            "commercial_terms": {
                "price_basis": "Delivered Hosur DC, per requested unit of measure",
                "tax": "GST extra, shown separately",
                "payment_terms": "60 days from invoice",
                "validity_days": 90,
                "freight": "Include freight in rates or state it separately",
            },
            "evaluation_rules": {
                "criteria": [
                    {"name": "Evaluated cost (normalized, ex-GST)", "weight": 60},
                    {"name": "Questionnaire / quality qualification", "weight": 25},
                    {"name": "Lead time and commercial terms", "weight": 15},
                ],
                "qualification": "Supplier must meet all must-pass questionnaire items (Q1-Q4).",
            },
        },
        "lines": [{"line_no": l[0], "description": l[1], "specification": l[2], "quantity": l[3], "uom": l[4], "category": l[5]} for l in LINES],
        "questions": [{"q_no": q[0], "question": q[1], "answer_type": q[2], "must_pass": q[3], "pass_rule": q[4], "requires_evidence": q[5]} for q in QUESTIONS],
        "suppliers": SUPPLIERS,
        "sample_responses": [
            {"supplier": "shree-balaji", "version": 1, "channel": "email", "received_at": "2026-10-29T11:42:00+05:30",
             "email_subject": "Quotation against RFQ-2026-0142", "email_body": "Please find attached our quotation. Regards, Venkatesh",
             "files": [{"path": f"responses/shree-balaji/{sb1}", "kind": "quote"}]},
            {"supplier": "deccan", "version": 1, "channel": "email", "received_at": "2026-10-30T16:05:00+05:30",
             "email_subject": "DPI quotation - RFQ-2026-0142", "email_body": "Dear team, our quotation and ISO certificate are attached. - Anita",
             "files": [{"path": f"responses/deccan/{dpi}", "kind": "quote"}, {"path": f"responses/deccan/{dpi_cert}", "kind": "attachment"}]},
            {"supplier": "pacific", "version": 1, "channel": "email", "received_at": "2026-10-31T08:20:00+05:30",
             "email_subject": "Pacific Pack Solutions - offer RFQ-2026-0142", "email_body": "Please see our offer letter attached. Best, Marcus",
             "files": [{"path": f"responses/pacific/{pps}", "kind": "quote"}]},
            {"supplier": "hosur-box", "version": 1, "channel": "whatsapp", "received_at": "2026-10-29T17:46:00+05:30",
             "email_subject": None, "email_body": "Sir rate card photo sending. pls check. - Murugan",
             "files": [{"path": f"responses/hosur-box/{hbw}", "kind": "quote"}]},
            {"supplier": "sri-lakshmi", "version": 1, "channel": "email", "received_at": "2026-11-02T09:14:00+05:30",
             "email_subject": "RE: RFQ-2026-0142 Corrugated packaging - quote", "email_body": None,
             "files": [{"path": f"responses/sri-lakshmi/{slp}", "kind": "email_body"}]},
            {"supplier": "shree-balaji", "version": 2, "channel": "email", "received_at": "2026-11-02T15:10:00+05:30",
             "email_subject": "REVISED quotation - RFQ-2026-0142",
             "email_body": "Sir, please consider our revised quotation. Item 3 rate corrected (typing error) and duplicate entry for master carton removed. Regards, Venkatesh",
             "files": [{"path": f"responses/shree-balaji/{sb2}", "kind": "quote"}], "is_correction": True},
        ],
    }
    with open(os.path.join(ROOT, "demo-rfq.json"), "w") as f:
        json.dump(manifest, f, indent=2)
    print("generated:", sb1, sb2, dpi, dpi_cert, pps, hbw, slp)


if __name__ == "__main__":
    main()
